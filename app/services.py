"""Общие финансовые и авторизационные правила для HTML и REST."""
import csv
import hashlib
import io
import re
import secrets
import sqlite3
from datetime import date, datetime, timezone
from functools import wraps

from flask import current_app, g, has_request_context
from werkzeug.security import generate_password_hash

from .db import atomic, get_db

ROLES = ('Super Admin', 'Admin', 'Cashier', 'Investor')
MAX_MONEY = 9223372036854775807
INVALID_TEXT = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


class DomainError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def _fresh_authorized_user(user):
    """Повторно проверить исполнителя внутри транзакции записи.

    Аутентификация маршрута проходит до вызова сервиса, поэтому параллельную
    блокировку, смену роли, сброс пароля или отзыв сессии/токена нужно проверить
    непосредственно перед чувствительным изменением.
    """
    if not isinstance(user, dict) or not user.get('id'):
        raise DomainError('Недостаточно прав для этого действия.', 403)
    row = get_db().execute('SELECT * FROM Users WHERE id=?', (user['id'],)).fetchone()
    if row is None or not row['is_active']:
        raise DomainError('Недостаточно прав для этого действия.', 403)
    if user.get('auth_version') is not None and user['auth_version'] != row['auth_version']:
        raise DomainError('Сессия больше недействительна.', 401)
    if has_request_context():
        now = datetime.now(timezone.utc).replace(tzinfo=None).isoformat(timespec='seconds')
        session_id = getattr(g, 'web_session_id', None)
        if session_id is not None:
            valid = get_db().execute('''SELECT 1 FROM WebSessions
                WHERE id=? AND user_id=? AND revoked_at IS NULL AND expires_at>? AND auth_version=?''',
                (session_id, row['id'], now, row['auth_version'])).fetchone()
            if valid is None:
                raise DomainError('Сессия больше недействительна.', 401)
        token_id = getattr(g, 'api_token_id', None)
        if token_id is not None:
            valid = get_db().execute('''SELECT 1 FROM ApiTokens
                WHERE id=? AND user_id=? AND active=1 AND expires_at>?''',
                (token_id, row['id'], now)).fetchone()
            if valid is None:
                raise DomainError('Токен недействителен или отозван.', 401)
    return dict(row)


def write_operation(function):
    @wraps(function)
    def wrapped(*args, **kwargs):
        try:
            with atomic():
                if has_request_context() and args and isinstance(args[0], dict):
                    args = (_fresh_authorized_user(args[0]),) + args[1:]
                return function(*args, **kwargs)
        except sqlite3.IntegrityError as error:
            raise DomainError('Запись уже существует или нарушает целостность данных.', 409) from error
    return wrapped


def require_role(user, *roles):
    """Проверить активность пользователя и его роль до бизнес-операции."""
    if not user or not user['is_active'] or user['role'] not in roles:
        raise DomainError('Недостаточно прав для этого действия.', 403)


def text_value(value, label, maximum, required=True):
    if not isinstance(value, str):
        raise DomainError(f'{label}: ожидается текст.')
    value = value.strip()
    if INVALID_TEXT.search(value):
        raise DomainError(f'{label}: недопустимые управляющие символы.')
    if (required and not value) or len(value) > maximum:
        raise DomainError(f'{label}: допустимо от {1 if required else 0} до {maximum} символов.')
    return value


def integer(value, label='ID'):
    if isinstance(value, bool) or not isinstance(value, int) or not 0 < value <= MAX_MONEY:
        raise DomainError(f'{label}: требуется положительное целое число.')
    return value


def identifier(value):
    if isinstance(value, str) and value.isascii() and value.isdigit():
        if len(value) > 19:
            raise DomainError('ID: число выходит за допустимый диапазон.')
        value = int(value)
    return integer(value)


def enum_value(value, allowed, label):
    if not isinstance(value, str) or value not in allowed:
        raise DomainError(f'{label}: недопустимое значение.')
    return value


def timestamp(value=None):
    if value is None or value == '':
        return datetime.now(timezone.utc).replace(tzinfo=None).isoformat(timespec='seconds')
    if not isinstance(value, str):
        raise DomainError('Дата должна быть в формате ISO 8601.')
    try:
        parsed = datetime.fromisoformat(value)
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
        return parsed.isoformat(timespec='seconds')
    except (ValueError, OverflowError) as error:
        raise DomainError('Некорректная дата операции.') from error


def password_hash(password):
    if not isinstance(password, str) or not 8 <= len(password) <= 256:
        raise DomainError('Пароль должен содержать от 8 до 256 символов.')
    try:
        password.encode('utf-8')
    except UnicodeEncodeError as error:
        raise DomainError('Пароль содержит недопустимые символы Unicode.') from error
    return generate_password_hash(password, method='pbkdf2:sha256:600000')


def record(table, record_id):
    # Имена таблиц — внутренние константы, а не значения из запроса.
    if table not in ('Users', 'Funds', 'Transactions', 'ApiTokens'):
        raise ValueError('Неизвестная таблица')
    # Единственный подставляемый идентификатор проверен по фиксированному списку выше.
    row = get_db().execute(f'SELECT * FROM {table} WHERE id=?', (identifier(record_id),)).fetchone()  # nosec B608
    if row is None:
        raise DomainError('Запись не найдена.', 404)
    return dict(row)


def allowed_fund_ids(user, global_view=False):
    """Вернуть только фонды, доступные роли и назначенным Rights пользователя."""
    require_role(user, *ROLES)
    db = get_db()
    if global_view:
        # Глобальная сводка разрешена только Super Admin и Investor и включает
        # исключительно фонды `for_stats`, даже если у Investor есть другие Rights.
        require_role(user, 'Super Admin', 'Investor')
        rows = db.execute("SELECT id FROM Funds WHERE type='for_stats'")
    elif user['role'] == 'Super Admin':
        rows = db.execute('SELECT id FROM Funds')
    elif user['role'] == 'Investor':
        rows = db.execute("SELECT f.id FROM Funds f JOIN Rights r ON r.fund_id=f.id WHERE r.user_id=? AND f.type='for_stats'", (user['id'],))
    else:
        rows = db.execute('SELECT fund_id AS id FROM Rights WHERE user_id=?', (user['id'],))
    return {row['id'] for row in rows}


def require_fund(user, fund_id, writing=False):
    """Проверить Rights к фонду и отдельно запретить запись в архив."""
    fund_id = identifier(fund_id)
    if fund_id not in allowed_fund_ids(user):
        raise DomainError('Нет доступа к этому фонду.', 403)
    fund = record('Funds', fund_id)
    if writing and not fund['is_active']:
        raise DomainError('Фонд архивирован. Создание и изменение операций недоступно.', 409)
    return fund


def ledger(fund_id):
    """Рассчитать текущий баланс фонда в целых минорных единицах.

    Доход увеличивает баланс, расход уменьшает его, а перевод учитывается по
    стороне пары: `debit` уменьшает источник, `credit` увеличивает получателя.
    """
    rows = get_db().execute('''SELECT * FROM Transactions WHERE
      (type='income' AND to_fund_id=?) OR (type='expense' AND from_fund_id=?) OR
      (type='inter-transaction' AND ((transfer_side='debit' AND from_fund_id=?) OR
      (transfer_side='credit' AND to_fund_id=?)))''', (fund_id,)*4).fetchall()
    balance = sum(row['money'] if row['type'] == 'income' or row['transfer_side'] == 'credit'
                  else -row['money'] for row in rows)
    return balance, len(rows)


def fund_info(fund):
    result = dict(fund)
    result['balance'], result['transaction_count'] = ledger(result['id'])
    return result


def list_funds(user, include_archived=True):
    ids = allowed_fund_ids(user)
    return [fund_info(row) for row in get_db().execute('SELECT * FROM Funds ORDER BY is_active DESC,name,id')
            if row['id'] in ids and (include_archived or row['is_active'])]


def fund_detail(user, fund_id):
    return fund_info(require_fund(user, fund_id))


def parse_filters(filters=None):
    filters = dict(filters or {})
    result = {}
    for field in ('date_from', 'date_to'):
        if filters.get(field):
            try:
                result[field] = date.fromisoformat(filters[field]).isoformat()
            except (ValueError, TypeError) as error:
                raise DomainError('Период: используйте даты YYYY-MM-DD.') from error
    if result.get('date_from', '') > result.get('date_to', '9999-12-31'):
        raise DomainError('Начало периода не может быть позже окончания.')
    if filters.get('type'):
        result['type'] = enum_value(filters['type'], ('income','expense','inter-transaction'), 'Тип операции')
    if filters.get('pay_type'):
        result['pay_type'] = text_value(filters['pay_type'], 'Способ оплаты', 50)
    if filters.get('fund_id'):
        result['fund_id'] = identifier(filters['fund_id'])
    return result


def _joined_transactions():
    return get_db().execute('''SELECT t.*, u.username, u.fullname, u.role AS author_role,
      ff.name AS from_fund_name, tf.name AS to_fund_name
      FROM Transactions t JOIN Users u ON u.id=t.user_id
      LEFT JOIN Funds ff ON ff.id=t.from_fund_id LEFT JOIN Funds tf ON tf.id=t.to_fund_id
      ORDER BY t.datetime DESC,t.id DESC''')


def list_transactions(user, filters=None, global_view=False):
    filters = parse_filters(filters)
    ids = allowed_fund_ids(user, global_view)
    if filters.get('fund_id') and filters['fund_id'] not in ids:
        raise DomainError('Нет доступа к этому фонду.', 403)
    selected = {filters['fund_id']} if filters.get('fund_id') else ids
    results, seen = [], set()
    for raw in _joined_transactions():
        row = dict(raw)
        if row['from_fund_id'] not in selected and row['to_fund_id'] not in selected:
            continue
        if filters.get('date_from') and row['datetime'][:10] < filters['date_from']:
            continue
        if filters.get('date_to') and row['datetime'][:10] > filters['date_to']:
            continue
        if any(filters.get(key) and filters[key] != row[key] for key in ('type', 'pay_type')):
            continue
        if row['transfer_id']:
            if row['transfer_id'] in seen:
                continue
            seen.add(row['transfer_id'])
        # Видимый журнал не должен раскрывать недоступную встречную сторону фонда.
        for prefix in ('from', 'to'):
            if row[f'{prefix}_fund_id'] not in ids:
                if row[f'{prefix}_fund_id'] is not None:
                    row[f'{prefix}_fund_name'] = 'Недоступный фонд'
                row[f'{prefix}_fund_id'] = None
        results.append(row)
    return results


def _transaction_data(user, data, transfer=False, allow_archived=False):
    if not isinstance(data, dict):
        raise DomainError('Ожидается объект с полями операции.')
    amount = integer(data.get('money'), 'Сумма в минорных единицах')
    kind = 'inter-transaction' if transfer else enum_value(data.get('type'), ('income', 'expense'), 'Тип операции')
    name = text_value(data.get('name') or {'income':'Приём оплаты','expense':'Расход','inter-transaction':'Межфондовый перевод'}[kind], 'Название', 50)
    description = text_value(data.get('description', ''), 'Описание', 150, False)
    payment = data.get('pay_type', '')
    if transfer and payment in ('', None):
        payment = 'Внутренний перевод'
    pay_type = text_value(payment, 'Способ оплаты', 50)
    if transfer:
        from_id, to_id = identifier(data.get('from_fund_id')), identifier(data.get('to_fund_id'))
        if from_id == to_id:
            raise DomainError('Выберите разные фонды для перевода.')
        require_fund(user, from_id, not allow_archived)
        require_fund(user, to_id, not allow_archived)
    else:
        fund_id = identifier(data.get('fund_id'))
        require_fund(user, fund_id, not allow_archived)
        from_id, to_id = (None, fund_id) if kind == 'income' else (fund_id, None)
    return dict(name=name, description=description, money=amount, type=kind, pay_type=pay_type,
                datetime=timestamp(data.get('datetime')), from_fund_id=from_id, to_fund_id=to_id)


def _insert_transaction(values, user_id, transfer_id=None, side=None):
    columns = ('name','description','money','type','pay_type','datetime','from_fund_id','to_fund_id')
    cursor = get_db().execute('''INSERT INTO Transactions
      (name,description,money,type,pay_type,datetime,from_fund_id,to_fund_id,user_id,transfer_id,transfer_side)
      VALUES(?,?,?,?,?,?,?,?,?,?,?)''', tuple(values[key] for key in columns)+(user_id,transfer_id,side))
    return cursor.lastrowid


@write_operation
def create_transaction(user, data):
    require_role(user, 'Super Admin', 'Admin', 'Cashier')
    values = _transaction_data(user, data)
    return record('Transactions', _insert_transaction(values, user['id']))


@write_operation
def create_transfer(user, data):
    """Атомарно создать две связанные записи межфондового перевода."""
    require_role(user, 'Super Admin', 'Admin')
    values = _transaction_data(user, data, True)
    transfer_id = secrets.token_hex(16)
    # Перевод перемещает деньги между уже существующими фондами и поэтому
    # не считается внешним доходом или расходом отчётности.
    # Обе стороны межфондового перевода сохраняются одной atomic-транзакцией:
    # ошибка списания или зачисления откатывает всю операцию.
    debit = _insert_transaction(values, user['id'], transfer_id, 'debit')
    credit = _insert_transaction(values, user['id'], transfer_id, 'credit')
    result = record('Transactions', debit)
    result['linked_transaction_id'] = credit
    return result


def _require_modify(user, transaction):
    require_role(user, 'Super Admin', 'Admin')
    for field in ('from_fund_id', 'to_fund_id'):
        if transaction[field] is not None:
            # Super Admin может исправлять архивную историю; остальные роли
            # меняют операции только пока все затронутые фонды активны.
            require_fund(user, transaction[field], writing=user['role'] != 'Super Admin')
    if user['role'] == 'Admin' and transaction['user_id'] != user['id']:
        author = record('Users', transaction['user_id'])
        if author['role'] != 'Cashier':
            raise DomainError('Можно исправлять только свои операции и операции кассиров.', 403)


def transaction_for_modify(user, transaction_id):
    """Вернуть доступную для изменения операцию без раскрытия чужих ID."""
    transaction = record('Transactions', transaction_id)
    try:
        _require_modify(user, transaction)
    except DomainError as error:
        if error.status == 403:
            raise DomainError('Запись не найдена.', 404) from error
        raise
    return transaction


def can_modify(user, transaction):
    try:
        _require_modify(user, record('Transactions', transaction['id']))
        return True
    except DomainError:
        return False


@write_operation
def edit_transaction(user, transaction_id, data):
    original = transaction_for_modify(user, transaction_id)
    merged = dict(original)
    merged['fund_id'] = original['to_fund_id'] if original['type'] == 'income' else original['from_fund_id']
    merged.update(data)
    is_transfer = original['type'] == 'inter-transaction'
    if is_transfer and merged.get('type') != 'inter-transaction':
        raise DomainError('Тип межфондового перевода нельзя изменить.')
    values = _transaction_data(user, merged, is_transfer, user['role'] == 'Super Admin')
    columns = ('name','description','money','type','pay_type','datetime','from_fund_id','to_fund_id')
    selector = 'transfer_id' if is_transfer else 'id'
    # Идентификаторы взяты из фиксированных констант; внешние значения остаются параметрами.
    get_db().execute(f"UPDATE Transactions SET {','.join(key+'=?' for key in columns)} WHERE {selector}=?",  # nosec B608
                     tuple(values[key] for key in columns)+(original['transfer_id'] if is_transfer else original['id'],))
    return record('Transactions', transaction_id)


@write_operation
def delete_transaction(user, transaction_id):
    original = transaction_for_modify(user, transaction_id)
    if original['transfer_id']:
        get_db().execute('DELETE FROM Transactions WHERE transfer_id=?', (original['transfer_id'],))
    else:
        get_db().execute('DELETE FROM Transactions WHERE id=?', (original['id'],))


def last_cashier_transaction_id(user):
    require_role(user, 'Cashier', 'Super Admin')
    row = get_db().execute('SELECT id FROM Transactions WHERE user_id=? ORDER BY id DESC LIMIT 1', (user['id'],)).fetchone()
    return row['id'] if row else None


@write_operation
def cancel_last(user, transaction_id):
    require_role(user, 'Cashier', 'Super Admin')
    original = record('Transactions', transaction_id)
    if original['user_id'] != user['id'] or original['id'] != last_cashier_transaction_id(user) or original['type'] == 'inter-transaction':
        raise DomainError('Запись не найдена.', 404)
    try:
        require_fund(user, original['to_fund_id'] or original['from_fund_id'],
                     writing=user['role'] != 'Super Admin')
    except DomainError as error:
        if error.status == 403:
            raise DomainError('Запись не найдена.', 404) from error
        raise
    get_db().execute('DELETE FROM Transactions WHERE id=?', (original['id'],))


def _public_user(row):
    return {key: row[key] for key in ('id','username','fullname','role','is_active')}


def list_users(user, search='', role=''):
    require_role(user, 'Super Admin')
    if role:
        enum_value(role, ROLES, 'Роль')
    rows = get_db().execute('SELECT * FROM Users ORDER BY id')
    return [_public_user(row) for row in rows if (not role or row['role'] == role)
            and (not search or search.casefold() in (row['username']+' '+row['fullname']).casefold())]


@write_operation
def create_user(user, data):
    require_role(user, 'Super Admin')
    username = text_value(data.get('username'), 'Логин', 50)
    fullname = text_value(data.get('fullname'), 'Полное имя', 150)
    role = enum_value(data.get('role'), ROLES, 'Роль')
    cursor = get_db().execute('INSERT INTO Users(username,fullname,password_hash,role) VALUES(?,?,?,?)',
                             (username,fullname,password_hash(data.get('password')),role))
    return _public_user(record('Users', cursor.lastrowid))


def _protect_superadmin_access(actor, target, *, role=None, is_active=None):
    """Не дать административному изменению удалить всех активных Super Admin."""
    next_role = target['role'] if role is None else role
    next_active = bool(target['is_active']) if is_active is None else is_active
    removes_active_superadmin = (
        target['role'] == 'Super Admin' and target['is_active']
        and (next_role != 'Super Admin' or not next_active)
    )
    if not removes_active_superadmin:
        return
    if actor['id'] == target['id']:
        if next_role != 'Super Admin':
            raise DomainError('Нельзя изменить роль собственной учётной записи Super Admin.', 409)
        raise DomainError('Нельзя заблокировать собственную учётную запись.', 409)
    remaining = get_db().execute("""SELECT COUNT(*) FROM Users
        WHERE id<>? AND role='Super Admin' AND is_active=1""", (target['id'],)).fetchone()[0]
    if remaining == 0:
        raise DomainError('В системе должен оставаться хотя бы один активный Super Admin.', 409)


@write_operation
def edit_user(user, user_id, data):
    require_role(user, 'Super Admin')
    original = record('Users', user_id)
    values = {**original, **data}
    role = enum_value(values['role'],ROLES,'Роль')
    _protect_superadmin_access(user, original, role=role)
    get_db().execute('UPDATE Users SET username=?,fullname=?,role=?,auth_version=auth_version+1 WHERE id=?',
                    (text_value(values['username'],'Логин',50), text_value(values['fullname'],'Полное имя',150),
                     role, original['id']))
    if 'password' in data:
        reset_password(user, original['id'], data['password'])
    return _public_user(record('Users', user_id))


@write_operation
def block_user(user, user_id, is_active):
    require_role(user, 'Super Admin')
    target = record('Users', user_id)
    if not isinstance(is_active, bool):
        raise DomainError('is_active должен быть true или false.')
    _protect_superadmin_access(user, target, is_active=is_active)
    get_db().execute('UPDATE Users SET is_active=?,auth_version=auth_version+1 WHERE id=?', (int(is_active),user_id))
    return _public_user(record('Users', user_id))


@write_operation
def reset_password(user, user_id, password):
    require_role(user, 'Super Admin')
    record('Users', user_id)
    get_db().execute('UPDATE Users SET password_hash=?,auth_version=auth_version+1 WHERE id=?', (password_hash(password),user_id))
    get_db().execute('UPDATE ApiTokens SET active=0 WHERE user_id=?', (user_id,))


@write_operation
def create_fund(user, data):
    require_role(user, 'Super Admin')
    cursor = get_db().execute('INSERT INTO Funds(name,description,type,user_id) VALUES(?,?,?,?)',
       (text_value(data.get('name'),'Название',50), text_value(data.get('description',''),'Описание',150,False),
        enum_value(data.get('type'),('for_stats','no_stats'),'Тип фонда'),user['id']))
    return fund_info(record('Funds', cursor.lastrowid))


@write_operation
def edit_fund(user, fund_id, data):
    require_role(user, 'Super Admin')
    original = record('Funds', fund_id)
    values = {**original, **data}
    get_db().execute('UPDATE Funds SET name=?,description=?,type=? WHERE id=?',
       (text_value(values['name'],'Название',50),text_value(values['description'],'Описание',150,False),
        enum_value(values['type'],('for_stats','no_stats'),'Тип фонда'),original['id']))
    return fund_info(record('Funds', fund_id))


@write_operation
def archive_fund(user, fund_id):
    require_role(user, 'Super Admin')
    record('Funds', fund_id)
    get_db().execute('UPDATE Funds SET is_active=0 WHERE id=?', (fund_id,))
    return fund_info(record('Funds', fund_id))


def list_rights(user):
    require_role(user, 'Super Admin')
    return [dict(row) for row in get_db().execute('SELECT * FROM Rights ORDER BY user_id,fund_id')]


@write_operation
def grant_right(user, user_id, fund_id):
    require_role(user, 'Super Admin')
    record('Users', user_id)
    record('Funds', fund_id)
    get_db().execute('INSERT OR IGNORE INTO Rights(user_id,fund_id) VALUES(?,?)', (user_id,fund_id))
    return dict(get_db().execute('SELECT * FROM Rights WHERE user_id=? AND fund_id=?',(user_id,fund_id)).fetchone())


@write_operation
def revoke_right(user, user_id, fund_id):
    require_role(user, 'Super Admin')
    get_db().execute('DELETE FROM Rights WHERE user_id=? AND fund_id=?', (identifier(user_id),identifier(fund_id)))


@write_operation
def set_rights(user, user_id, fund_ids):
    require_role(user, 'Super Admin')
    record('Users', user_id)
    ids = {identifier(fund_id) for fund_id in fund_ids}
    for fund_id in ids:
        record('Funds', fund_id)
    get_db().execute('DELETE FROM Rights WHERE user_id=?',(user_id,))
    get_db().executemany('INSERT INTO Rights(user_id,fund_id) VALUES(?,?)',[(user_id,fund_id) for fund_id in ids])


def list_tokens(user):
    require_role(user, 'Super Admin')
    return [dict(row) for row in get_db().execute('''SELECT t.id,t.user_id,t.active,t.datetime,t.expires_at,u.username
      FROM ApiTokens t JOIN Users u ON u.id=t.user_id ORDER BY t.id DESC''')]


@write_operation
def create_token(user, user_id):
    """Создать Bearer-токен и вернуть секрет ровно один раз.

    В базе сохраняется только SHA-256; срок действия задаёт
    `API_TOKEN_LIFETIME`, а отзыв владельца проверяется на каждом API-запросе.
    """
    require_role(user, 'Super Admin')
    owner = record('Users', user_id)
    if not owner['is_active']:
        raise DomainError('Пользователь заблокирован.',409)
    token = secrets.token_hex(64)
    created = datetime.now(timezone.utc).replace(tzinfo=None)
    expires = created + current_app.config['API_TOKEN_LIFETIME']
    cursor = get_db().execute('INSERT INTO ApiTokens(token,datetime,expires_at,user_id) VALUES(?,?,?,?)',
                             (hashlib.sha256(token.encode()).hexdigest(),
                              created.isoformat(timespec='seconds'), expires.isoformat(timespec='seconds'), owner['id']))
    return {'id':cursor.lastrowid,'token':token,'user_id':owner['id']}


@write_operation
def revoke_token(user, token_id):
    require_role(user, 'Super Admin')
    record('ApiTokens',token_id)
    get_db().execute('UPDATE ApiTokens SET active=0 WHERE id=?',(token_id,))


def list_web_sessions(user):
    require_role(user, 'Super Admin')
    now = timestamp()
    return [dict(row) for row in get_db().execute('''SELECT s.id,s.user_id,s.created_at,s.last_seen_at,
        s.expires_at,s.ip_address,s.user_agent,u.username,u.fullname,u.role
        FROM WebSessions s JOIN Users u ON u.id=s.user_id
        WHERE s.revoked_at IS NULL AND s.expires_at>? AND s.auth_version=u.auth_version AND u.is_active=1
        ORDER BY s.last_seen_at DESC,s.id DESC''', (now,))]


@write_operation
def revoke_web_session(user, session_id):
    require_role(user, 'Super Admin')
    session_id = identifier(session_id)
    row = get_db().execute('SELECT id,revoked_at,expires_at FROM WebSessions WHERE id=?',
                           (session_id,)).fetchone()
    if row is None:
        raise DomainError('Сессия не найдена.', 404)
    if row['revoked_at'] is not None or row['expires_at'] <= timestamp():
        raise DomainError('Сессия уже завершена.', 409)
    get_db().execute('UPDATE WebSessions SET revoked_at=?,revoked_by=? WHERE id=?',
                     (timestamp(), user['id'], session_id))
    return session_id


@write_operation
def revoke_all_web_sessions(user, keep_session_id=None):
    require_role(user, 'Super Admin')
    now = timestamp()
    if keep_session_id is None:
        cursor = get_db().execute('''UPDATE WebSessions SET revoked_at=?,revoked_by=?
            WHERE revoked_at IS NULL AND expires_at>?''', (now, user['id'], now))
    else:
        keep_session_id = identifier(keep_session_id)
        cursor = get_db().execute('''UPDATE WebSessions SET revoked_at=?,revoked_by=?
            WHERE revoked_at IS NULL AND expires_at>? AND id<>?''',
            (now, user['id'], now, keep_session_id))
    return cursor.rowcount


def dashboard(user, filters=None):
    """Собрать сводку только по доступным `for_stats` фондам за выбранный период."""
    require_role(user, 'Super Admin', 'Investor')
    filters = parse_filters(filters)
    # Период дашборда меняет только потоки; баланс остаётся текущим балансом журнала.
    transactions = list_transactions(user, filters, True)
    funds = [fund_info(row) for row in get_db().execute("SELECT * FROM Funds WHERE type='for_stats' ORDER BY name,id")]
    for fund in funds:
        fund['income'] = sum(row['money'] for row in transactions if row['type']=='income' and row['to_fund_id']==fund['id'])
        fund['expense'] = sum(row['money'] for row in transactions if row['type']=='expense' and row['from_fund_id']==fund['id'])
    income = sum(row['money'] for row in transactions if row['type']=='income')
    expense = sum(row['money'] for row in transactions if row['type']=='expense')
    month = datetime.now(timezone.utc).strftime('%Y-%m')
    month_rows = list_transactions(user, {'date_from':month+'-01','date_to':datetime.now(timezone.utc).date().isoformat()}, True)
    month_income = sum(row['money'] for row in month_rows if row['type']=='income')
    month_expense = sum(row['money'] for row in month_rows if row['type']=='expense')
    daily = {}
    for row in transactions:
        if row['type']=='income':
            day = row['datetime'][:10]
            daily[day] = daily.get(day,0)+row['money']
    return dict(balance=sum(fund['balance'] for fund in funds), turnover=income, expenses=expense,
                net_flow=income-expense, month_income=month_income, month_expense=month_expense,
                profit=month_income-month_expense, funds=funds, recent=transactions[:10],
                trend=[{'date':day,'money':daily[day]} for day in sorted(daily)],
                distribution=[{'name':fund['name'],'balance':fund['balance']} for fund in funds])


def export_rows(user, filters=None):
    require_role(user, 'Super Admin', 'Investor')
    keys = ('id','datetime','name','description','type','money','pay_type','from_fund_id','from_fund_name',
            'to_fund_id','to_fund_name','fullname')
    return [{key:row[key] for key in keys} for row in list_transactions(user, filters, True)]


def export_file(user, filters, format):
    """Сформировать CSV или XLSX отчёт с теми же фильтрами и ограничениями доступа."""
    enum_value(format,('csv','xlsx'),'Формат отчета')
    rows = export_rows(user, filters)
    columns = ('id','datetime','name','description','type','money','pay_type','from_fund_id','from_fund_name','to_fund_id','to_fund_name','fullname')
    headers = ('ID','Дата (UTC)','Название','Описание','Тип','Сумма (minor units)','Способ оплаты','ID фонда-источника','Фонд-источник','ID фонда-получателя','Фонд-получатель','Автор')
    if format == 'csv':
        def safe(value):
            if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')):
                return "'"+value
            return value
        output = io.StringIO(newline='')
        writer = csv.writer(output,delimiter=';')
        writer.writerow(headers)
        writer.writerows([safe(row[key]) for key in columns] for row in rows)
        return output.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8','cybran-report.csv'
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Операции for_stats'
    sheet.append(headers)
    for row in rows:
        # Excel сохраняет только 15 значащих цифр; большие целые записываем текстом.
        sheet.append([str(row[key]) if isinstance(row[key],int) and abs(row[key]) >= 10**15 else row[key] for key in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):
                cell.data_type = 's'  # Пользовательский текст не должен стать формулой Excel.
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF',bold=True)
        cell.fill = PatternFill('solid',fgColor='B4232D')
    for index in range(1,len(columns)+1):
        sheet.column_dimensions[get_column_letter(index)].width = 24 if index>1 else 10
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','cybran-report.xlsx'
