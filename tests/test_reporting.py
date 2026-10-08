"""Сверка общей отчётности с внешними потоками и реальными балансами фондов."""
import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app import services as svc


PERIOD = {"date_from": "2026-04-01", "date_to": "2026-04-30"}


@pytest.fixture
def reporting_data(seed_transaction, api):
    visible = [seed_transaction(money=10000, fund_id=1, name="Visible income"),
               seed_transaction(money=3000, fund_id=2, name="Global income without detail rights"),
               seed_transaction(money=2000, fund_id=1, kind="expense", name="Visible expense")]
    seed_transaction(money=999999, fund_id=3, username="root", name="HIDDEN NO STATS OPERATION")
    for source, target, amount in [(1, 2, 1500), (3, 1, 600), (1, 3, 100)]:
        response = api("POST", "/api/transactions/add_transfer", "root", {
            "from_fund_id": source, "to_fund_id": target, "money": amount,
            "datetime": "2026-04-16T09:00:00", "name": "Internal transfer"})
        assert response.status_code == 201
    return visible


@pytest.mark.parametrize("username", ["root", "investor"])
def test_dashboard_kpis_and_balances_exclude_internal_flows(app, users, reporting_data, username):
    with app.app_context():
        result = svc.dashboard(users[username], PERIOD)
    assert result["balance"] == 11500
    assert result["turnover"] == 13000
    assert result["expenses"] == 2000
    assert result["net_flow"] == 11000
    assert type(result["balance"]) is type(result["turnover"]) is int
    funds = {fund["id"]: fund for fund in result["funds"]}
    assert set(funds) == {1, 2, 4}
    assert (funds[1]["balance"], funds[2]["balance"]) == (7000, 4500)
    assert funds[1]["income"] == 10000 and funds[1]["expense"] == 2000
    assert funds[2]["income"] == 3000 and funds[2]["expense"] == 0
    assert result["trend"] == [{"date": "2026-04-15", "money": 13000}]
    assert len(result["recent"]) == 6
    assert [row["datetime"] for row in result["recent"]] == sorted(
        [row["datetime"] for row in result["recent"]], reverse=True)
    assert "HIDDEN NO STATS" not in str(result)


def test_dashboard_period_filters_flows_not_current_balance(app, users, reporting_data, seed_transaction):
    seed_transaction(money=700, day="2026-03-31T23:59:59")
    seed_transaction(money=300, day="2026-05-01T00:00:00")
    with app.app_context():
        result = svc.dashboard(users["investor"], PERIOD)
    assert result["balance"] == 12500
    assert result["turnover"] == 13000 and result["expenses"] == 2000
    assert len(result["recent"]) == 6


def test_month_kpis_are_independent_from_chosen_period(app, users, seed_transaction):
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    seed_transaction(money=711, day=f"{today}T00:00:00")
    seed_transaction(money=211, kind="expense", day=f"{today}T00:00:00")
    seed_transaction(money=99999, fund_id=3, day=f"{today}T00:00:00")
    with app.app_context():
        result = svc.dashboard(users["investor"], {"date_from": "2020-01-01", "date_to": "2020-01-31"})
    assert result["turnover"] == result["expenses"] == 0
    assert result["month_income"] == 711 and result["month_expense"] == 211
    assert result["profit"] == 500


def test_report_scope_includes_all_stats_but_redacts_hidden_counterpart(app, users, reporting_data):
    with app.app_context():
        rows = svc.export_rows(users["investor"], PERIOD)
    assert len(rows) == 6
    assert any(row["name"] == "Global income without detail rights" for row in rows)
    assert "HIDDEN NO STATS" not in str(rows)
    assert all(row["from_fund_id"] != 3 and row["to_fund_id"] != 3 for row in rows)
    assert any(row["type"] == "inter-transaction" and row["from_fund_id"] is None for row in rows)
    assert any(row["type"] == "inter-transaction" and row["to_fund_id"] is None for row in rows)
    assert "token" not in str(rows) and "password_hash" not in str(rows)


@pytest.mark.parametrize("username", ["root", "investor"])
def test_csv_and_real_excel_export_match(client, login, reporting_data, username):
    assert login(username).status_code == 302
    csv_response = client.get("/reports/export", query_string={**PERIOD, "format": "csv"})
    assert csv_response.status_code == 200
    assert "attachment" in csv_response.headers["Content-Disposition"]
    assert "text/csv" in csv_response.content_type
    csv_rows = list(csv.reader(StringIO(csv_response.data.decode("utf-8-sig")), delimiter=";"))
    assert len(csv_rows) == 7
    assert len(csv_rows[0]) == 12
    assert sum(int(row[5]) for row in csv_rows[1:] if row[4] == "income") == 13000
    assert "HIDDEN NO STATS" not in str(csv_rows)

    xlsx_response = client.get("/reports/export", query_string={**PERIOD, "format": "xlsx"})
    assert xlsx_response.status_code == 200
    assert "spreadsheetml" in xlsx_response.content_type
    workbook = load_workbook(BytesIO(xlsx_response.data), data_only=False)
    sheet = workbook.active
    xlsx_rows = list(sheet.values)
    assert len(xlsx_rows) == len(csv_rows)
    assert [["" if value is None else str(value) for value in row] for row in xlsx_rows] == csv_rows
    assert all(type(row[5]) is int for row in xlsx_rows[1:])


def test_export_date_boundaries_and_invalid_requests(client, login, seed_transaction):
    chosen = seed_transaction(day="2026-04-30T23:59:59")
    seed_transaction(day="2026-05-01T00:00:00")
    assert login("investor").status_code == 302
    response = client.get("/reports/export", query_string={**PERIOD, "format": "csv"})
    records = list(csv.reader(StringIO(response.data.decode("utf-8-sig")), delimiter=";"))
    assert len(records) == 2 and int(records[1][0]) == chosen
    assert client.get("/reports/export?format=pdf").status_code == 400
    assert client.get("/reports/export?format=csv&date_from=bad").status_code == 400
    assert client.get("/reports/export?format=csv&fund_id=3").status_code == 403


def test_untrusted_text_is_not_a_csv_or_excel_formula(client, login, seed_transaction):
    seed_transaction(name="=1+1", description="@SUM(1,1)")
    assert login("root").status_code == 302
    response = client.get("/reports/export?format=csv")
    values = list(csv.reader(StringIO(response.data.decode("utf-8-sig")), delimiter=";"))[1]
    assert values[2] == "'=1+1" and values[3] == "'@SUM(1,1)"
    response = client.get("/reports/export?format=xlsx")
    sheet = load_workbook(BytesIO(response.data), data_only=False).active
    assert sheet["C2"].value == "=1+1" and sheet["C2"].data_type == "s"
    assert sheet["D2"].value == "@SUM(1,1)" and sheet["D2"].data_type == "s"


def test_fund_reporting_switch_immediately_changes_dashboard_and_export(app, users, seed_transaction):
    seed_transaction(money=5000)
    with app.app_context():
        assert svc.dashboard(users["investor"])["balance"] == 5000
        svc.edit_fund(users["root"], 1, {"type": "no_stats"})
        assert svc.dashboard(users["investor"])["balance"] == 0
        assert svc.export_rows(users["investor"]) == []
        with pytest.raises(svc.DomainError) as error:
            svc.fund_detail(users["investor"], 1)
        assert error.value.status == 403
        svc.edit_fund(users["root"], 1, {"type": "for_stats"})
        assert svc.dashboard(users["investor"])["balance"] == 5000


def test_zero_and_negative_distribution_render_without_corrupting_balance(client, login, seed_transaction):
    assert login("investor").status_code == 302
    assert client.get("/dashboard").status_code == 200
    seed_transaction(kind="expense", money=5000)
    response = client.get("/dashboard")
    assert response.status_code == 200
    text = response.get_data(as_text=True)
    assert "−50,00" in text
    assert "положительные остатки" in text


def test_excel_preserves_minor_unit_integer_beyond_15_digits(client, login, seed_transaction):
    exact = 9223372036854775807
    seed_transaction(money=exact)
    assert login("root").status_code == 302
    csv_response = client.get("/reports/export?format=csv")
    csv_rows = list(csv.reader(StringIO(csv_response.data.decode("utf-8-sig")), delimiter=";"))
    assert csv_rows[1][5] == str(exact)
    xlsx_response = client.get("/reports/export?format=xlsx")
    cell = load_workbook(BytesIO(xlsx_response.data), data_only=False).active["F2"]
    assert cell.value == str(exact) and cell.data_type == "s"
